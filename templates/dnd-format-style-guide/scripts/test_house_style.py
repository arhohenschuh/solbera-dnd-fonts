"""Regression checks for the bundled House Style v3 builder and validator."""

import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

import build_adventure_docx as builder
import validate_adventure_docx as checker


class HouseStyleV3Tests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.output = Path(self.directory.name) / "node.docx"

    def build(self, blocks=None):
        document = builder.new_document()
        builder.render(document, builder.DEMO if blocks is None else blocks)
        return document

    def findings(self, document, node=None):
        document.save(self.output)
        return checker.validate(self.output, node)[0]

    def paragraph(self, document, style):
        return next(paragraph for paragraph in document.paragraphs if paragraph.style.name == style)

    def assert_detected(self, document, message, level="CRITICAL"):
        findings = self.findings(document)
        self.assertTrue(any(severity == level and message in text for severity, text in findings), findings)

    def test_demo_passes_without_post_build_restyling(self):
        self.assertEqual([], self.findings(self.build()))

    def test_example_and_node_selection_pass(self):
        example = Path(builder.HERE).parent / "references" / "example_node.json"
        blocks = json.loads(example.read_text(encoding="utf-8"))["_blocks"]
        self.assertEqual([], self.findings(self.build(blocks), "N17"))

    def test_named_styles_and_numbering_are_preserved(self):
        self.build().save(self.output)
        with ZipFile(builder.TEMPLATE) as template, ZipFile(self.output) as generated:
            self.assertEqual(template.read("word/styles.xml"), generated.read("word/styles.xml"))
            original = etree.fromstring(template.read("word/numbering.xml"))
            result = etree.fromstring(generated.read("word/numbering.xml"))
            self.assertEqual(etree.tostring(original, method="c14n"), etree.tostring(result, method="c14n"))

    def test_sample_content_is_not_rendered(self):
        document = self.build()
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for placeholder in ("Title Card", "Credits", "Someone wrote an epigraph", "Paraphrase me"):
            self.assertNotIn(placeholder, text)
        self.assertEqual("N00: Sample Node", document.paragraphs[0].text)
        self.assertEqual(1, len(document.sections))

    def test_template_cannot_be_silently_replaced_by_blank_document(self):
        with patch.object(builder, "TEMPLATE", str(Path(self.directory.name) / "missing.docx")):
            with self.assertRaises(FileNotFoundError):
                builder.new_document()

    def test_inherited_v3_size_is_resolved(self):
        document = self.build()
        self.assertIsNone(document.styles["Core Bulleted"].font.size)
        self.assertEqual(10, checker._style_property(document.styles["Core Bulleted"], "size"))
        self.assertEqual([], self.findings(document))

    def test_v1_fonts_are_rejected_in_named_styles(self):
        for style, legacy in (("Heading 2", "Andada SC"), ("Core Body", "TeX Gyre Bonum")):
            with self.subTest(style=style):
                document = self.build()
                document.styles[style].font.name = legacy
                self.assert_detected(document, f"Style '{style}'")

    def test_required_unused_style_is_checked(self):
        document = self.build()
        document.styles["Heading 4"].font.size = Pt(16)
        self.assert_detected(document, "Style 'Heading 4': size")

    def test_legacy_font_and_wrong_v3_role_are_rejected_in_runs(self):
        for font in ("TeX Gyre Bonum", "Scaly Sans Remake"):
            with self.subTest(font=font):
                document = self.build()
                self.paragraph(document, "Core Body").runs[0].font.name = font
                self.assert_detected(document, "name must be 'Bookinsanity Remake'")

    def test_linked_character_font_override_is_checked(self):
        document = self.build()
        run = self.paragraph(document, "Heading 2").runs[0]
        linked_id = document.styles["Heading 2"].element.find(qn("w:link")).get(qn("w:val"))
        linked = next(style for style in document.styles if style.style_id == linked_id)
        linked.font.name = "Andada SC"
        run.style = linked
        self.assert_detected(document, "name must be 'Baskervville Caps'")

    def test_heading_size_override_is_checked(self):
        document = self.build()
        self.paragraph(document, "Heading 2").runs[0].font.size = Pt(22)
        self.assert_detected(document, "size must be 18.5")

    def test_synthetic_caps_are_rejected(self):
        for attribute in ("small_caps", "all_caps"):
            with self.subTest(attribute=attribute):
                document = self.build()
                setattr(self.paragraph(document, "Heading 3").runs[0].font, attribute, True)
                self.assert_detected(document, f"synthetic {attribute}")

    def test_inherited_accent_override_is_checked(self):
        document = self.build()
        document.styles["Heading 3"].font.color.rgb = RGBColor.from_string("FF0000")
        self.assert_detected(document, "color must be '920000'")

    def test_box_border_is_checked(self):
        document = self.build()
        border = document.styles["Boxed Text"].element.find(f"{qn('w:pPr')}/{qn('w:pBdr')}/{qn('w:top')}")
        border.set(qn("w:color"), "FF0000")
        self.assert_detected(document, "incorrect top border", "MAJOR")

    def test_tuning_label_color_is_checked(self):
        document = self.build()
        self.paragraph(document, "Sidebar Bulleted").runs[0].font.color.rgb = RGBColor.from_string("000000")
        self.assert_detected(document, "label is not maroon", "MAJOR")

    def test_table_typography_checks_every_cell(self):
        for row, font in ((0, "Scaly Sans Remake"), (2, "Bookinsanity Remake")):
            with self.subTest(row=row):
                document = self.build()
                run = next(run for run in document.tables[0].cell(row, 2).paragraphs[0].runs if run.text.strip())
                run.font.name = font
                self.assert_detected(document, f"Table 1, row {row + 1}, cell 3: name")

    def test_table_size_is_checked(self):
        document = self.build()
        run = next(run for run in document.tables[0].cell(1, 1).paragraphs[0].runs if run.text.strip())
        run.font.size = Pt(10)
        self.assert_detected(document, "size must be 9")

    def test_table_banding_borders_and_alignment_are_checked(self):
        document = self.build()
        cell = document.tables[0].cell(2, 2)
        cell._tc.get_or_add_tcPr().find(qn("w:shd")).set(qn("w:fill"), "FFFFFF")
        cell._tc.get_or_add_tcPr().find(qn("w:tcBorders")).remove(
            cell._tc.get_or_add_tcPr().find(f"{qn('w:tcBorders')}/{qn('w:right')}")
        )
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        findings = self.findings(document)
        for expected in ("direct fill must be F2DBDB", "right border must be", "vertically centered"):
            self.assertTrue(any(expected in text for _, text in findings), findings)

    def test_footer_font_and_page_field_are_checked(self):
        document = self.build()
        footer = document.sections[0].footer
        footer.paragraphs[0].runs[0].font.name = "Bookinsanity Remake"
        for field in list(footer._element.iter(qn("w:instrText"))):
            field.getparent().remove(field)
        findings = self.findings(document)
        for expected in ("Running footer: name", "needs a PAGE field"):
            self.assertTrue(any(expected in text for _, text in findings), findings)

    def test_rebuilding_footer_does_not_duplicate_it(self):
        document = self.build()
        builder._set_running_footer(document, "N00: Sample Node")
        self.assertEqual(1, len(document.sections[0].footer.paragraphs))
        self.assertEqual([], self.findings(document))

    def test_tuning_block_keeps_its_four_tiers_together(self):
        document = self.build()
        headline = self.paragraph(document, "Sidebar Headline")
        tiers = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Sidebar Bulleted"]
        self.assertTrue(headline.paragraph_format.keep_with_next)
        self.assertEqual([True, True, True, False], [paragraph.paragraph_format.keep_with_next for paragraph in tiers])
        self.assertTrue(all(paragraph.paragraph_format.keep_together for paragraph in tiers))

    def test_incomplete_tuning_is_not_silently_rendered(self):
        with self.assertRaisesRegex(ValueError, "Tuning requires exactly"):
            self.build([{"type": "tuning", "tiers": {"Baseline": "Default"}}])

    def test_saved_document_reopens(self):
        self.build().save(self.output)
        reopened = Document(self.output)
        self.assertEqual("N00: Sample Node", reopened.paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()
